"""Build the LogicScript self-checking expense tracker (works in Google Sheets, Excel, LibreOffice)."""
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = 'content/expense-tracker-lesson/LogicScript_Expense_Tracker.xlsx'
FIRST, LAST = 6, 505          # 500 expense rows
F = 'Arial'
INK, MUTED = '1F2937', '6B7280'
INPUT_FILL = PatternFill('solid', fgColor='FFF9DB')
HEAD_FILL = PatternFill('solid', fgColor='0B1021')
RULE_FILL = PatternFill('solid', fgColor='EEF2F7')
RED = PatternFill('solid', fgColor='FDE2E1'); AMBER = PatternFill('solid', fgColor='FEF3C7'); GREEN = PatternFill('solid', fgColor='DCFCE7')
thin = Side(style='thin', color='D1D5DB'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)


def font(**kw):
    kw.setdefault('name', F); kw.setdefault('color', INK)
    return Font(**kw)


wb = Workbook()

# ---------------- Settings ----------------
st = wb.active; st.title = 'Settings'
st['A1'] = 'Settings'; st['A1'].font = font(size=16, bold=True)
st['A2'] = 'Yellow cells are yours to change. Everything else updates automatically.'; st['A2'].font = font(color=MUTED, italic=True)
settings = [
    ('Tax year', 2026, 'The year this tracker covers. Dates outside it are flagged.'),
    ('Large expense: review at or above', 1000, 'Your own review rule (example value). Flags, never blocks.'),
    ('Receipt required at or above', 75, 'Your own record-keeping rule (example value), not a legal threshold. Ask your tax professional what records you need.'),
    ('Currency label', 'USD', 'Shown in headers only; no conversion is done.'),
]
for i, (label, value, note) in enumerate(settings, start=4):
    st.cell(i, 1, label).font = font(bold=True)
    c = st.cell(i, 2, value); c.font = font(color='0000FF'); c.fill = INPUT_FILL; c.border = BOX
    st.cell(i, 3, note).font = font(color=MUTED)
st['A10'] = 'Categories'; st['A10'].font = font(bold=True, size=12)
st['C10'] = 'Generic examples. Rename or add your own (up to 30). Map them to your tax forms with a tax professional.'; st['C10'].font = font(color=MUTED)
cats = ['Advertising & marketing', 'Bank & payment fees', 'Equipment', 'Meals', 'Office supplies',
        'Phone & internet', 'Professional services', 'Rent', 'Software & subscriptions', 'Training & education',
        'Travel', 'Vehicle', 'Other']
for i in range(30):
    c = st.cell(11 + i, 2, cats[i] if i < len(cats) else None); c.fill = INPUT_FILL; c.border = BOX; c.font = font(color='0000FF')
st['A43'] = 'Payment methods'; st['A43'].font = font(bold=True, size=12)
for i, m in enumerate(['Card', 'Bank transfer', 'Cash', 'Check', 'Other']):
    c = st.cell(44 + i, 2, m); c.fill = INPUT_FILL; c.border = BOX; c.font = font(color='0000FF')
st.column_dimensions['A'].width = 34; st.column_dimensions['B'].width = 28; st.column_dimensions['C'].width = 90
CATS = 'Settings!$B$11:$B$40'; YEAR = 'Settings!$B$4'; BIG = 'Settings!$B$5'; RCPT = 'Settings!$B$6'

# ---------------- Expenses ----------------
ex = wb.create_sheet('Expenses', 0)
ex['A1'] = 'Expenses'; ex['A1'].font = font(size=18, bold=True)
ex['A2'] = ('Type in the yellow columns. Each grey "rule" column checks one thing. Status: ERROR = left out of totals until fixed; '
            'CHECK = counted, review it, then set Reviewed? to Yes; OK = passed.')
ex['A2'].font = font(color=MUTED, italic=True)
ex['A3'] = 'Rows marked EXAMPLE contain deliberate mistakes so you can see the checks work. Delete them before you start.'
ex['A3'].font = font(color='B45309', bold=True)
headers = ['Date', 'Vendor', 'Description', 'Category', 'Amount', 'Payment method', 'Receipt saved?', 'Notes',
           'Reviewed?', 'Status', 'Issues', 'Rule: date', 'Rule: amount', 'Rule: category', 'Rule: duplicate', 'Rule: receipt', 'Rule: size']
for col, h in enumerate(headers, start=1):
    c = ex.cell(5, col, h); c.font = font(bold=True, color='FFFFFF'); c.fill = HEAD_FILL
    c.alignment = Alignment(vertical='center', wrap_text=True); c.border = BOX
ex.row_dimensions[5].height = 30
widths = [11, 18, 18, 23, 10, 12, 9, 9, 11, 9, 41, 24, 24, 20, 20, 18, 20]
for col, w in enumerate(widths, start=1):
    ex.column_dimensions[ex.cell(5, col).column_letter].width = w

R = f'${FIRST}:$'  # unused helper
rng = lambda col: f'${col}${FIRST}:${col}${LAST}'
for r in range(FIRST, LAST + 1):
    for col in range(1, 10):
        c = ex.cell(r, col); c.fill = INPUT_FILL; c.font = font(); c.border = BOX
    ex.cell(r, 1).number_format = 'yyyy-mm-dd'
    ex.cell(r, 5).number_format = '#,##0.00;[Red]-#,##0.00'
    has = f'COUNTA($A{r}:$H{r})>0'   # Reviewed? alone does not make a row
    rules = {
        12: f'=IF(NOT({has}),"",IF(A{r}="","Missing date. ",IF(NOT(ISNUMBER(A{r})),"Date is not a real date. ",'
            f'IF(YEAR(A{r})<>{YEAR},"Date outside tax year. ",IF(A{r}>TODAY(),"Date is in the future. ","")))))',
        13: f'=IF(NOT({has}),"",IF(E{r}="","Missing amount. ",IF(NOT(ISNUMBER(E{r})),"Amount is not a number. ",'
            f'IF(E{r}<=0,"Amount must be above zero. ",""))))',
        14: f'=IF(NOT({has}),"",IF(D{r}="","Missing category. ",IF(COUNTIF({CATS},D{r})=0,"Category not in your list. ","")))',
        15: f'=IF(AND(ISNUMBER(A{r}),ISNUMBER(E{r}),B{r}<>""),IF(COUNTIFS({rng("A")},A{r},{rng("B")},B{r},{rng("E")},E{r})>1,'
            f'"Possible duplicate. ",""),"")',
        16: f'=IF(AND(ISNUMBER(E{r}),E{r}>={RCPT},G{r}<>"Yes"),"Receipt not marked as saved. ","")',
        17: f'=IF(AND(ISNUMBER(E{r}),E{r}>={BIG}),"Large amount: review. ","")',
    }
    for col, formula in rules.items():
        c = ex.cell(r, col, formula); c.fill = RULE_FILL; c.font = font(color=MUTED, size=9); c.border = BOX
    # Errors can never be cleared by review; warnings clear when the owner marks the row reviewed.
    s = ex.cell(r, 10, f'=IF(NOT({has}),"",IF(L{r}&M{r}&N{r}<>"","ERROR",IF(AND(O{r}&P{r}&Q{r}<>"",I{r}<>"Yes"),"CHECK","OK")))')
    s.font = font(bold=True); s.border = BOX; s.alignment = Alignment(horizontal='center')
    i = ex.cell(r, 11, f'=TRIM(L{r}&M{r}&N{r}&O{r}&P{r}&Q{r})&IF(AND(I{r}="Yes",O{r}&P{r}&Q{r}<>"",L{r}&M{r}&N{r}=""),"  (reviewed)","")')
    i.font = font(size=11); i.border = BOX

examples = [
    (date(2026, 1, 5), 'Office supply store', 'Printer paper and toner', 'Office supplies', 48.20, 'Card', 'Yes'),
    (date(2026, 1, 12), 'Design software co.', 'Monthly subscription', 'Software & subscriptions', 59.99, 'Card', 'Yes'),
    (date(2026, 1, 12), 'Design software co.', 'Monthly subscription', 'Software & subscriptions', 59.99, 'Card', 'Yes'),
    (date(2026, 2, 3), 'Airline', 'Flight to client meeting', 'Travel', 412.30, 'Card', ''),
    (date(2026, 2, 14), 'Corner cafe', 'Client lunch', 'Meal', 38.75, 'Card', 'Yes'),
    (date(2026, 3, 1), 'Internet provider', 'March internet', 'Phone & internet', -89.99, 'Bank transfer', 'Yes'),
    (date(2025, 12, 30), 'Computer store', 'Laptop', 'Equipment', 1299.00, 'Card', 'Yes'),
    (date(2026, 3, 15), 'Office landlord', 'March office rent', 'Rent', 1200.00, 'Bank transfer', 'Yes'),
]
for k, row in enumerate(examples):
    r = FIRST + k
    for col, v in enumerate(row, start=1):
        if v != '':
            ex.cell(r, col, v)
    ex.cell(r, 8, 'EXAMPLE - delete')

dv_date = DataValidation(type='date', operator='greaterThan', formula1='DATE(2000,1,1)', allow_blank=True,
                         showErrorMessage=True, errorTitle='Date', error='Enter a real date, e.g. 2026-03-15.')
dv_amt = DataValidation(type='decimal', operator='greaterThan', formula1='0', allow_blank=True,
                        showErrorMessage=True, errorTitle='Amount', error='Enter a positive number without currency symbols.')
dv_cat = DataValidation(type='list', formula1=f'={CATS}', allow_blank=True, showErrorMessage=True,
                        errorTitle='Category', error='Pick a category from the list (edit the list on the Settings tab).')
dv_pay = DataValidation(type='list', formula1='=Settings!$B$44:$B$48', allow_blank=True, showErrorMessage=True)
dv_rc = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True, showErrorMessage=True)
dv_rv = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True, showErrorMessage=True)
for dv, col in ((dv_date, 'A'), (dv_amt, 'E'), (dv_cat, 'D'), (dv_pay, 'F'), (dv_rc, 'G'), (dv_rv, 'I')):
    ex.add_data_validation(dv); dv.add(f'{col}{FIRST}:{col}{LAST}')
status_rng = f'J{FIRST}:J{LAST}'
ex.conditional_formatting.add(status_rng, CellIsRule(operator='equal', formula=['"ERROR"'], fill=RED, font=Font(name=F, bold=True, color='B91C1C')))
ex.conditional_formatting.add(status_rng, CellIsRule(operator='equal', formula=['"CHECK"'], fill=AMBER, font=Font(name=F, bold=True, color='B45309')))
ex.conditional_formatting.add(status_rng, CellIsRule(operator='equal', formula=['"OK"'], fill=GREEN, font=Font(name=F, bold=True, color='15803D')))
ex.conditional_formatting.add(f'A{FIRST}:I{LAST}', FormulaRule(formula=[f'$J{FIRST}="ERROR"'], fill=RED))
ex.freeze_panes = 'A6'
ex['J5'].comment = Comment('ERROR rows are excluded from all totals until fixed. CHECK rows are counted; mark Reviewed? = Yes once you have checked them.', 'LogicScript')
ex['I5'].comment = Comment('Set to Yes after you have checked a CHECK row (e.g. a real large expense). Errors cannot be cleared this way.', 'LogicScript')
ex['L4'] = 'Rule columns: do not type here'; ex['L4'].font = font(color=MUTED, italic=True, size=9)

# ---------------- Summary ----------------
sm = wb.create_sheet('Summary', 1)
sm['A1'] = 'Summary & checks'; sm['A1'].font = font(size=18, bold=True)
sm['A2'] = '=" Tax year "&Settings!$B$4&"  |  amounts in "&Settings!$B$7'; sm['A2'].font = font(color=MUTED)
xr = lambda col: 'Expenses!' + rng(col)   # Summary must point at the Expenses tab
st_r = xr('J'); amt = xr('E')
panel = [
    ('Overall', None),
    ('Rows entered', f'=COUNTIF({st_r},"OK")+COUNTIF({st_r},"CHECK")+COUNTIF({st_r},"ERROR")'),
    ('Rows with ERRORS (left out of totals)', f'=COUNTIF({st_r},"ERROR")'),
    ('Rows to CHECK (counted, not yet reviewed)', f'=COUNTIF({st_r},"CHECK")'),
    ('Rows OK', f'=COUNTIF({st_r},"OK")'),
    ('Possible duplicates', f'=COUNTIF({xr("O")},"Possible*")'),
    ('Receipts not marked as saved', f'=COUNTIF({xr("P")},"Receipt*")'),
    ('Amount left out of totals (rows with errors)', f'=SUMIFS({amt},{st_r},"ERROR")'),
    ('Total counted (OK + CHECK)', f'=SUMIFS({amt},{st_r},"OK")+SUMIFS({amt},{st_r},"CHECK")'),
    ('Category totals add up to total?', '=IF(ROUND(SUM($B$24:$B$53)-B12,2)=0,"YES","NO")'),
    ('Monthly totals add up to total?', '=IF(ROUND(SUM($E$24:$E$35)-B12,2)=0,"YES","NO")'),
]
for i, (label, formula) in enumerate(panel, start=4):
    sm.cell(i, 1, label).font = font(bold=(formula is None))
    if formula:
        c = sm.cell(i, 2, formula); c.font = font(bold=True); c.border = BOX
sm['B4'] = ('=IF(B5=0,"Add your first expense",IF(AND(B6=0,B7=0,B13="YES",B14="YES"),"ALL CHECKS PASSED",'
            'B6&" error(s), "&B7&" to review"))')
sm['B4'].font = font(bold=True, size=13); sm['B4'].border = BOX
for cell in ('B11', 'B12'):
    sm[cell].number_format = '#,##0.00'
sm.conditional_formatting.add('B4', FormulaRule(formula=['$B$4="ALL CHECKS PASSED"'], fill=GREEN, font=Font(name=F, bold=True, color='15803D')))
sm.conditional_formatting.add('B4', FormulaRule(formula=['$B$6>0'], fill=RED, font=Font(name=F, bold=True, color='B91C1C')))
sm.conditional_formatting.add('B4', FormulaRule(formula=['AND($B$6=0,$B$7>0)'], fill=AMBER, font=Font(name=F, bold=True, color='B45309')))
sm.conditional_formatting.add('B13:B14', CellIsRule(operator='equal', formula=['"NO"'], fill=RED))

sm['A22'] = 'By category'; sm['A22'].font = font(bold=True, size=12)
sm['D22'] = 'By month'; sm['D22'].font = font(bold=True, size=12)
for c, h in (('A23', 'Category'), ('B23', 'Total counted'), ('D23', 'Month'), ('E23', 'Total counted')):
    sm[c] = h; sm[c].font = font(bold=True, color='FFFFFF'); sm[c].fill = HEAD_FILL
for i in range(30):
    r = 24 + i
    sm.cell(r, 1, f'=IF(Settings!$B${11 + i}="","",Settings!$B${11 + i})')
    v = sm.cell(r, 2, f'=IF(A{r}="","",SUMIFS({amt},{xr("D")},A{r},{st_r},"OK")+SUMIFS({amt},{xr("D")},A{r},{st_r},"CHECK"))')
    v.number_format = '#,##0.00'
months = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']
for m, name in enumerate(months, start=1):
    r = 23 + m
    sm.cell(r, 4, name)
    lo = f'DATE({YEAR},{m},1)'; hi = f'DATE({YEAR},{m + 1},1)'
    v = sm.cell(r, 5, '=' + '+'.join(
        f'SUMIFS({amt},{xr("A")},">="&{lo},{xr("A")},"<"&{hi},{st_r},"{s}")' for s in ('OK', 'CHECK')))
    v.number_format = '#,##0.00'
sm.column_dimensions['A'].width = 44; sm.column_dimensions['B'].width = 22; sm.column_dimensions['D'].width = 14; sm.column_dimensions['E'].width = 16

# ---------------- Start Here ----------------
hs = wb.create_sheet('Start Here', 0)
lines = [
    ('LogicScript Expense Tracker', 'title'),
    ('A free tracker that checks its own entries. Made for freelancers and small businesses.', 'sub'),
    ('', None),
    ('How to use it', 'h'),
    ('1. Settings tab: set your tax year, your two review rules, and your categories (yellow cells).', None),
    ('2. Expenses tab: delete the rows marked EXAMPLE, then type one expense per row in the yellow columns.', None),
    ('3. Watch the Status column: ERROR = fix it (left out of totals), CHECK = counted but review it, OK = passed.', None),
    ('   After reviewing a CHECK row (for example a real large expense), set Reviewed? to Yes. Errors cannot be cleared that way.', None),
    ('4. Summary tab: you are done when it says ALL CHECKS PASSED.', None),
    ('', None),
    ('What the rules check', 'h'),
    ('Date: missing, not a real date, outside your tax year, or in the future.', None),
    ('Amount: missing, not a number, or zero/negative (refunds: record them separately in Notes).', None),
    ('Category: missing or not in your list (catches typos that slip past the dropdown when pasting).', None),
    ('Duplicate: same date, vendor and amount entered more than once.', None),
    ('Receipt: at or above your receipt rule but not marked as saved.', None),
    ('Size: at or above your large-expense rule, so you look at it twice.', None),
    ('', None),
    ('Why both dropdowns AND rule columns?', 'h'),
    ('Dropdowns stop typos when you type, but pasted or imported values can skip them. The rule columns re-check every row anyway.', None),
    ('', None),
    ('Important', 'h'),
    ('This is a record-keeping tool, not tax, legal or accounting advice. The example thresholds and categories are placeholders.', None),
    ('Ask a qualified tax professional which expenses are deductible and which records you must keep.', None),
    ('Works in Google Sheets (File > Make a copy), Microsoft Excel and LibreOffice. Your data stays in your own copy.', None),
]
for i, (text, kind) in enumerate(lines, start=1):
    c = hs.cell(i, 1, text)
    c.font = {'title': font(size=20, bold=True), 'sub': font(color=MUTED, size=12), 'h': font(bold=True, size=13)}.get(kind, font(size=11))
hs.column_dimensions['A'].width = 120
wb.active = 0
wb.save(OUT)
print('saved', OUT)
